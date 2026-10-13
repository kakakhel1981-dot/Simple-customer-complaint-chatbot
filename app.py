
import io
import re
import os
from pathlib import Path

import streamlit as st
from pypdf import PdfReader
from docx import Document

st.set_page_config(
    page_title="FLM Customer Support Assistant",
    page_icon="💬",
    layout="centered"
)

st.title("💬 FLM Customer Support Assistant")
st.caption("Your FLM procedure assistant")
st.write("How can I help you today?")


# --------------------------------------------------
# DOCUMENT READER
# --------------------------------------------------

def read_document(filename, file_bytes):
    ext = Path(filename).suffix.lower()

    if ext == ".txt":
        return file_bytes.decode("utf-8", errors="replace")

    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if ext == ".docx":
        doc = Document(io.BytesIO(file_bytes))
        lines = [
            p.text.strip()
            for p in doc.paragraphs
            if p.text.strip()
        ]

        for table in doc.tables:
            for row in table.rows:
                lines.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

        return "\n".join(lines)

    raise ValueError("Supported formats: PDF, DOCX, TXT")


def split_sections(text):
    # Keep paragraphs and numbered steps together where possible.
    paragraphs = [
        x.strip()
        for x in re.split(r"\n+", text)
        if x.strip()
    ]

    sections = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) > 1200:
            if current:
                sections.append(current)
                current = ""

            sections.extend(
                paragraph[i:i + 1000]
                for i in range(0, len(paragraph), 1000)
            )

        elif len(current) + len(paragraph) < 1200:
            current = (
                current + "\n" + paragraph
                if current else paragraph
            )

        else:
            sections.append(current)
            current = paragraph

    if current:
        sections.append(current)

    return sections


# --------------------------------------------------
# DOCUMENT SEARCH
# --------------------------------------------------

def normalise(text):
    return re.sub(
        r"[^a-z0-9\s]",
        " ",
        text.lower()
    ).strip()


def find_matching_sections(question, sections, limit=4):
    q = normalise(question)

    # Expand common support terms for retrieval only.
    replacements = {
        "sim mod": "sim modification",
        "sim modification": "sim modification",
        "apps": "applications",
        "application": "applications",
        "fix steps": "fixture steps",
        "fixation steps": "fixture steps",
    }

    for old, new in replacements.items():
        q = q.replace(old, new)

    stop_words = {
        "the", "is", "are", "of", "for", "to", "in",
        "a", "an", "and", "please", "tell", "me",
        "give", "show", "what", "how", "can", "you",
        "about", "i", "want", "know", "only", "explain"
    }

    terms = {
        word for word in q.split()
        if len(word) > 1 and word not in stop_words
    }

    if not terms:
        return []

    scored = []

    for index, section in enumerate(sections):
        content = normalise(section)

        score = sum(
            content.count(term)
            for term in terms
        )

        # Prefer sections that contain the full complaint name.
        if "sim modification" in q and "sim modification" in content:
            score += 5

        if score:
            scored.append((score, index, section))

    scored.sort(key=lambda x: (-x[0], x[1]))

    return [
        item[2]
        for item in scored[:limit]
    ]


def classify_question(question):
    q = normalise(question)

    if any(x in q for x in [
        "application", "applications", "apps",
        "system involved", "systems involved", "node involved"
    ]):
        return "applications"

    if any(x in q for x in [
        "fixture", "fixature", "fixation",
        "fix steps", "troubleshoot", "troubleshooting"
    ]):
        return "fixture"

    if any(x in q for x in [
        "procedure", "process", "steps",
        "how to", "how do i", "workflow"
    ]):
        return "procedure"

    if any(x in q for x in [
        "meaning", "define", "definition",
        "what is", "explain", "overview"
    ]):
        return "meaning"

    # A complaint name by itself requests a brief overview.
    return "overview"


def select_query(question, intent, complaint):
    if intent == "applications":
        return f"{complaint} applications involved systems applications nodes"

    if intent == "procedure":
        return f"{complaint} procedure process steps workflow"

    if intent == "fixture":
        return f"{complaint} fixture fixature fixation troubleshooting steps"

    if intent == "meaning":
        return f"{complaint} meaning definition overview description"

    if intent == "overview":
        return f"{complaint} meaning definition overview description"

    return f"{complaint} {question}"


def format_answer(question, sections, complaint):
    intent = classify_question(question)

    search_query = select_query(
        question,
        intent,
        complaint
    )

    matches = find_matching_sections(
        search_query,
        sections,
        limit=4
    )

    if not matches:
        return (
            "I could not find this information in the provided "
            "FLM document. Please try another keyword or check "
            "with your FLM document owner."
        )

    # For the complaint name or meaning, keep the response brief.
    if intent in ("overview", "meaning"):
        text = "\n".join(matches[:2])
        lines = [
            line.strip()
            for line in text.splitlines()
            if line.strip()
        ]

        # Three lines maximum, without creating new facts.
        return "\n".join(lines[:3])

    # For applications, procedure and fixture steps, show only
    # the sections relevant to the requested subject.
    return "\n\n".join(matches)


# --------------------------------------------------
# LOAD THE FLM DOCUMENT
# --------------------------------------------------

with st.sidebar:
    st.subheader("FLM Document")

    uploaded_file = st.file_uploader(
        "Upload approved FLM document",
        type=["pdf", "docx", "txt"]
    )

    st.caption(
        "Only this document is used as the knowledge source."
    )

document_bytes = None
document_name = None

if uploaded_file is not None:
    document_bytes = uploaded_file.getvalue()
    document_name = uploaded_file.name

elif os.getenv("FLM_DOC_PATH"):
    path = Path(os.getenv("FLM_DOC_PATH"))

    if path.is_file():
        document_bytes = path.read_bytes()
        document_name = path.name

sections = []

if document_bytes:
    try:
        text = read_document(
            document_name,
            document_bytes
        )
        sections = split_sections(text)

        if not sections:
            st.warning(
                "No readable text was found in the document."
            )

    except Exception as error:
        st.error(f"Could not read the FLM document: {error}")


# --------------------------------------------------
# CHAT INTERFACE
# --------------------------------------------------

if "messages" not in st.session_state:
    st.session_state.messages = []

if "complaint_name" not in st.session_state:
    st.session_state.complaint_name = "SIM Modification"

if not sections:
    st.info(
        "Please upload your approved FLM document in the sidebar "
        "to start chatting."
    )

# Show current in-memory conversation.
for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

question = st.chat_input(
    "Ask about SIM Modification, applications, procedure..."
)

if question:

    if not sections:
        st.warning("Upload an FLM document before asking questions.")

    else:
        # Remember the complaint topic temporarily.
        q = normalise(question)

        if "sim modification" in q or "sim mod" in q:
            st.session_state.complaint_name = "SIM Modification"

        elif "order deletion" in q:
            st.session_state.complaint_name = "Order Deletion"

        elif "provisioning" in q:
            st.session_state.complaint_name = "Provisioning"

        st.session_state.messages.append({
            "role": "user",
            "content": question
        })

        with st.chat_message("user"):
            st.markdown(question)

        answer = format_answer(
            question,
            sections,
            st.session_state.complaint_name
        )

        with st.chat_message("assistant"):
            st.markdown(answer)

        st.session_state.messages.append({
            "role": "assistant",
            "content": answer
        })

st.sidebar.divider()

if st.sidebar.button("Clear conversation"):
    st.session_state.messages = []
    st.session_state.complaint_name = "SIM Modification"
    st.rerun()

st.caption(
    "Privacy: the app does not deliberately write conversation "
    "history to a file or database. Messages are kept in Streamlit "
    "session state while the session is active. Hosting logs and "
    "platform-level data handling are separate."
)
